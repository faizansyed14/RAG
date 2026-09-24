from docx import Document as DocxDocument

from app.ingestion.docx_ingest import extract_blocks


def _make_docx(tmp_path):
    doc = DocxDocument()
    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph("This is the first paragraph of the document.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Drawing No"
    table.cell(0, 1).text = "Revision"
    table.cell(1, 0).text = "A-108"
    table.cell(1, 1).text = "A"
    doc.add_paragraph("This paragraph comes after the table.")
    path = str(tmp_path / "test.docx")
    doc.save(path)
    return path


def test_extract_blocks_preserves_order(tmp_path):
    path = _make_docx(tmp_path)
    blocks = extract_blocks(path)

    types = [b["type"] for b in blocks]
    assert types == ["heading", "paragraph", "table", "paragraph"]
    assert blocks[0]["text"] == "Project Overview"
    assert blocks[2]["rows"] == [["Drawing No", "Revision"], ["A-108", "A"]]
    assert blocks[3]["text"] == "This paragraph comes after the table."
