"""
DOCX ingestion. The tree engine's local mode only accepts PDF input (see
app/rag_service.py's module docstring), so this extracts the document's
real content -- paragraphs, headings, and tables kept as actual grid
tables, in true reading order -- renders it to a PDF (text_to_pdf.py), and
submits that.

python-docx's high-level `document.paragraphs` / `document.tables` lists
each type separately, losing interleaving order; walking
`document.element.body` directly (its children are `<w:p>` paragraph and
`<w:tbl>` table elements, in document order) and wrapping each back into
the matching python-docx object is what keeps a table appearing between
the two paragraphs it actually sits between, not shoved to the end.
"""

from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

_HEADING_STYLES = {"Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5"}


def _iter_body_blocks(document: DocxDocument):
    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield DocxParagraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, document)


def extract_blocks(file_path: str) -> list[dict]:
    document = DocxDocument(file_path)
    blocks: list[dict] = []

    for item in _iter_body_blocks(document):
        if isinstance(item, DocxParagraph):
            text = item.text.strip()
            if not text:
                continue
            is_heading = bool(item.style and item.style.name in _HEADING_STYLES)
            blocks.append({"type": "heading" if is_heading else "paragraph", "text": text})
        elif isinstance(item, DocxTable):
            rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
            rows = [row for row in rows if any(row)]
            if rows:
                blocks.append({"type": "table", "rows": rows})

    return blocks


def parse_docx(file_path: str, document_id) -> dict:
    from app.ingestion.blocks_ingest import submit_blocks

    blocks = extract_blocks(file_path)
    return submit_blocks(blocks, file_path, document_id)
