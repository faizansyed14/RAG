"""
XLSX ingestion. The tree engine's local mode only accepts PDF input (see
app/rag_service.py's module docstring), so this reads each non-empty sheet
into a heading + grid table block, in workbook order, and renders them
through the shared block pipeline (blocks_ingest.py). Formulas are read as
their last-computed value (data_only=True), not their expressions.
"""

from openpyxl import load_workbook


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def extract_blocks(file_path: str) -> list[dict]:
    workbook = load_workbook(file_path, data_only=True, read_only=True)
    try:
        blocks: list[dict] = []
        for sheet in workbook.worksheets:
            rows = [[_cell_str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(cell.strip() for cell in row)]
            if not rows:
                continue
            blocks.append({"type": "heading", "text": sheet.title})
            blocks.append({"type": "table", "rows": rows})
        return blocks
    finally:
        workbook.close()


def parse_xlsx(file_path: str, document_id) -> dict:
    from app.ingestion.blocks_ingest import submit_blocks

    blocks = extract_blocks(file_path)
    return submit_blocks(blocks, file_path, document_id)
